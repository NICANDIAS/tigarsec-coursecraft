"""
Generates lesson PowerPoints from a spec content area, following a specific pedagogical
structure (retrieval-practice / Rosenshine-style, not just "explain the topic"):

  1. Title
  2. Retrieval starter - recall questions about the PREVIOUS lesson (not this one) to
     activate prior learning before new content starts
  3. Learning objectives for THIS lesson
  4. Content, taught through guiding QUESTIONS posed to students rather than stated facts
     where possible (student-led), with real-world images and interspersed check questions
  5. Real-life example slides tying the content to an actual scenario
  6. Timed activities (and practicals, if enabled) - time allocated per activity, not fixed
  7. A retrieval plenary - checking THIS lesson's content, not a generic recap
  8. Objectives revisited - a self-assessment check against each objective

Two stages: (1) an LLM call producing a structured JSON outline, (2) a pure-Python
python-pptx renderer that turns that outline into an actual, visually varied deck.
Images resolve through three tiers - see resolve_slide_image.
"""
import os
import re
import uuid
import random
import requests
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

from .llm_gateway import complete_json

LESSON_OUTLINE_SYSTEM_PROMPT = """You are an experienced UK further-education lecturer's teaching \
assistant, building a lesson PowerPoint outline. Follow this pedagogical structure exactly - it is \
based on retrieval practice and student-led teaching, not a lecture-style content dump:

1. RETRIEVAL STARTER: 2-3 short recall questions about the PREVIOUS lesson's content (given below \
as context) - NOT about today's new topic. This activates prior learning before new content starts. \
If no previous lesson context is given, ask 2-3 foundational recall questions about concepts this \
new topic builds on instead.

2. LEARNING OBJECTIVES: 2-4 objectives for THIS lesson, phrased "By the end of this lesson, you \
will be able to...".

3. CONTENT: Break the topic into small segments. For EACH segment, write a GUIDING QUESTION to pose \
to students first (student-led - they attempt to answer BEFORE being told), then the key points as \
what the teacher confirms/adds afterward, then a separate CHECK QUESTION to verify understanding \
before moving on. Suggest a real-world image for each (image_query - a short descriptive phrase \
for finding a relevant photo, e.g. "network server room technician").

4. REAL-LIFE EXAMPLE(S): at least one slide per major segment linking the content to an actual \
real-world scenario relevant to the topic and industry, with a discussion question for students.

5. ACTIVITIES: student tasks (discussion, pair work, short exercise) with clear instructions, a \
grouping (pairs/individual/small group/whole class), and a sensible time allocation in minutes - \
base the time on the complexity of the task, don't default to a fixed number.

6. PRACTICALS (only if practicals_enabled is true): hands-on exercises appropriate to the subject \
(e.g., a Packet Tracer task for networking content), with instructions, materials needed, and a \
time allocation.

7. RETRIEVAL PLENARY: 2-3 questions checking what was taught in THIS lesson specifically (not \
generic), for use as an exit-ticket style check at the end.

8. OBJECTIVES REVIEW: restate each objective from step 2 with a short self-assessment prompt \
("Thumbs up if you can now...").

The sum of all activity/practical time_minutes plus a reasonable estimate for content delivery \
should roughly fit within the total lesson duration given - don't wildly overshoot it.

DURATION SELF-ASSESSMENT - do this BEFORE building the slides: given the actual depth of the \
selected content (how many learning outcomes/criteria, how much genuine understanding each requires) \
against the requested lesson duration, honestly assess whether this content can be properly taught - \
with real understanding, not a rushed skim - within that time. As a practical ceiling, a single \
class session tops out around 480 minutes (8 hours) - the longest a lecturer can request; if the \
content would need more than that, or more than the specific duration given, to do it justice, it \
does NOT fit.

LONG SESSION STRUCTURE - for any session longer than about 90 minutes, do NOT structure it as one \
long block of content followed by activities at the end. Alternate clearly: a chunk of content/ \
questioning, then a hands-on activity, practical, or short research task that lets students apply or \
explore what was just covered, before moving to the next content segment. For a genuinely long \
session (several hours), this means MORE activity/practical/research slides distributed THROUGHOUT \
the lesson, not one or two saved for the end - students should never sit through more than roughly \
30-40 minutes of direct content/questioning before a hands-on, research, or discussion break, \
regardless of overall lesson length. A research task can be a genuine activity slide too (grouping \
"pairs" or "small group", instructions to investigate/find out something specific and feed back) - \
not every break needs to be a practical.

If it does NOT fit: do not silently cram everything in, rush shallow coverage of every point, or \
skip content without saying so. Instead, scope THIS lesson to a coherent, complete subset of the \
content that genuinely fits the given duration (e.g. the foundational half of the topic), and set \
"fits_in_requested_time" to false with a "scope_note" explaining specifically what this lesson \
covers, what's left for a follow-up session, and a rough estimate of how much extra time that \
follow-up would need. The lecturer must be told clearly, never left to discover it themselves.

If it DOES fit: set "fits_in_requested_time" to true and leave "scope_note" as an empty string.

REVISITING WHAT THE CLASS FOUND HARD - if the previous lesson context contains "RESULTS FROM LAST \
LESSON'S LIVE CHECKS", those are real, anonymous results from the class's quiz and confidence poll, \
and the class did NOT secure those points. Then: (a) the retrieval_starter questions must target \
exactly those points (reworded - not the same question again); (b) straight after the objectives \
slide, add ONE content slide (two at most, only if the gaps are on unrelated topics) titled \
"Revisit: <short topic>", whose guiding_question probes the misunderstanding and whose key_points \
correct it explicitly - including why the most common wrong answer is wrong - with a check_question \
to confirm it is now secure; (c) then teach this lesson's new content as normal. Keep the revisit \
to about 10 minutes so the new content still fits the time.

RESOURCES - this tool must be a genuine one-stop shop: if an activity or practical describes
something the teacher needs to physically hand out or print (a sample data table, a set of cards to
cut up, a scenario brief, a worksheet, a fictional log file, a set of example emails) to actually RUN
it, that resource must be GENERATED here, not just referenced. Set the slide's "resource" to:
{"needed": true, "title": "string", "content": "string - the ACTUAL printable content, written out in
full with real (even if fictional) data - e.g. a real formatted routing table with genuine-looking
entries, or the actual text of 6 'evidence handling' statement cards, one per line, ready to print and
cut up"}. If an activity/practical needs nothing printed (e.g. a class discussion, or work done
entirely in a lab environment with no handout), set "resource" to null - do not invent a resource that
isn't actually needed.

CONCRETE PRACTICAL STEPS - an activity or practical's "instructions" must be genuinely actionable, \
not a vague description of the topic, and must be a LIST of separate step strings - one action per \
list entry - NOT one long string with the numbers written inline (e.g. NOT "1. Open a terminal. 2. \
Run: nmap -sV 192.168.1.0/24 3. ..." as a single string - that renders as an unreadable wall of text). \
Instead: ["Open a terminal.", "Run: nmap -sV 192.168.1.0/24", "Note which ports appear open on each \
host.", ...] - each entry is one discrete step, rendered as its own numbered line. If the practical \
involves a specific tool (e.g. Nmap, Wireshark, a specific OS command, a specific IDE), give the \
ACTUAL steps and commands to run, one per list entry, not just "use Nmap to scan the network". A \
lecturer or student should be able to follow the instructions directly without having to work out \
what to actually type or click themselves. A simple activity with no real sequential steps (e.g. "as \
a pair, discuss X") can be a single-entry list - the point is genuine steps never get squashed \
together into one paragraph, not that every activity needs many steps.

Return ONLY raw JSON (no markdown fences, no commentary) matching exactly this shape:
{
  "lesson_title": "string",
  "fits_in_requested_time": true,
  "scope_note": "string - empty if fits_in_requested_time is true",
  "slides": [
    {"type": "title", "title": "string", "subtitle": "string"},
    {"type": "retrieval_starter", "title": "Do Now: Last Lesson Recap", "questions": ["string", ...], "time_minutes": 5},
    {"type": "objectives", "objectives": ["string", ...]},
    {"type": "content", "title": "string", "guiding_question": "string", "key_points": ["string", ...], "check_question": "string", "image_query": "string"},
    {"type": "real_life_example", "title": "string", "scenario": "string", "discussion_question": "string", "image_query": "string"},
    {"type": "activity", "title": "string", "instructions": ["string - one step per entry", ...], "grouping": "pairs|individual|small group|whole class", "time_minutes": 0, "resource": null},
    {"type": "practical", "title": "string", "instructions": ["string - one step per entry", ...], "materials_needed": "string", "time_minutes": 0, "resource": null},
    {"type": "retrieval_plenary", "title": "Check Your Understanding", "questions": ["string", ...]},
    {"type": "objectives_review", "objectives_check": [{"objective": "string", "self_assess_prompt": "string"}]}
  ]
}

Include multiple content/real_life_example/activity slides as appropriate for the duration - a \
60-minute lesson should have more content coverage than a 20-minute one. Always include exactly \
one title, one retrieval_starter, one objectives, one retrieval_plenary, and one objectives_review \
slide, with as many content/real_life_example/activity/practical slides in between as make sense."""


def _estimate_outline_max_tokens(duration_minutes: int, include_practicals: bool) -> int:
    """
    Standing fix: a flat max_tokens=6000 genuinely truncated real lesson outlines mid-response
    (confirmed via a real "Unterminated string" error, and via response.stop_reason=='max_tokens'
    in server logs) - this single call has to hold every slide's full text AND, per the RESOURCES
    instruction above, the complete printable content of any practical/activity resource (a whole
    worksheet's worth of text, not a reference to one). Longer lessons need more content slides;
    practicals add substantial resource text on top. Scaled rather than a bigger flat guess, since
    a 20-minute lesson doesn't need the same budget as an 8-hour one.

    The outer cap was raised from 32000 to 48000 alongside the lesson-duration ceiling going from
    240 minutes to 8 hours - the old cap was already being reached even at the OLD 240-minute
    maximum (240*120+8000=36800), so a much longer lesson with more content/activity slides
    spread across it is even more likely to need the headroom. The model's own actual maximum
    output token limit isn't verified here (that's configured in llm_gateway.py, not reviewed
    this round) - if a very long lesson still truncates, that's the next thing to check.
    """
    base = 8000
    duration_component = duration_minutes * 120  # roughly accounts for more content/activity slides
    practicals_component = 6000 if include_practicals else 0
    return min(48000, base + duration_component + practicals_component)


def generate_lesson_outline(
    content_area_text: str,
    duration_minutes: int,
    previous_lesson_context: str,
    include_practicals: bool,
) -> dict:
    prompt = (
        f"Lesson duration: {duration_minutes} minutes\n"
        f"practicals_enabled: {include_practicals}\n"
        f"Previous lesson context: {previous_lesson_context or 'Not provided - use foundational recall questions instead.'}\n\n"
        f"Content area / spec criteria to build this lesson from:\n{content_area_text}"
    )
    max_tokens = _estimate_outline_max_tokens(duration_minutes, include_practicals)
    return complete_json(system=LESSON_OUTLINE_SYSTEM_PROMPT, prompt=prompt, max_tokens=max_tokens, for_draft_generation=True)


# ---------------------------------------------------------------------------
# Image resolution - three tiers, each optional, always falls through safely
# ---------------------------------------------------------------------------

def test_unsplash_connection() -> dict:
    """
    Makes a real, visible test call - unlike _fetch_unsplash_image (used during actual
    generation), this does NOT swallow errors, since the whole point is to show you what's
    actually going wrong instead of just falling back silently. Called from a button on the
    Media Library page.
    """
    access_key = os.getenv("UNSPLASH_ACCESS_KEY", "").strip()
    if not access_key:
        return {"ok": False, "message": "UNSPLASH_ACCESS_KEY is not set (or is empty) in .env."}

    try:
        resp = requests.get(
            "https://api.unsplash.com/search/photos",
            params={"query": "classroom", "per_page": 1},
            headers={"Authorization": f"Client-ID {access_key}"},
            timeout=10,
        )
    except requests.exceptions.RequestException as e:
        return {"ok": False, "message": f"Network error reaching Unsplash: {e}"}

    if resp.status_code == 401:
        return {
            "ok": False,
            "message": (
                "401 Unauthorized. Unsplash gives you two keys per app - an Access Key and a "
                "Secret Key. This integration needs the ACCESS KEY specifically. If you pasted "
                "the Secret Key by mistake, that's almost certainly this error - go back to "
                "unsplash.com/oauth/applications, open your app, and copy the Access Key instead."
            ),
        }
    if resp.status_code == 403:
        return {"ok": False, "message": f"403 Forbidden - Unsplash response: {resp.text[:300]}"}
    if resp.status_code == 429:
        return {"ok": False, "message": "429 Rate limited - Demo Unsplash apps are capped at 50 requests/hour."}
    if not resp.ok:
        return {"ok": False, "message": f"HTTP {resp.status_code} from Unsplash: {resp.text[:300]}"}

    results = resp.json().get("results", [])
    if not results:
        return {"ok": True, "message": "Connected successfully, but the test search returned no results (unusual, but not an auth problem)."}

    return {"ok": True, "message": f"Connected successfully - found {len(results)} result(s) for a test search."}


def _fetch_unsplash_image(query: str, cache_dir: str) -> tuple[str, str] | None:
    """Tier 1: a real photo via Unsplash, if UNSPLASH_ACCESS_KEY is configured. Returns
    (file_path, attribution_text) or None if unavailable/no results/any error - never raises,
    since a missing photo should just fall through to the next tier, not break generation."""
    access_key = os.getenv("UNSPLASH_ACCESS_KEY", "").strip()
    if not access_key:
        return None
    try:
        resp = requests.get(
            "https://api.unsplash.com/search/photos",
            params={"query": query, "per_page": 1, "orientation": "landscape"},
            headers={"Authorization": f"Client-ID {access_key}"},
            timeout=10,
        )
        resp.raise_for_status()
        results = resp.json().get("results", [])
        if not results:
            return None
        photo = results[0]
        image_url = photo["urls"]["regular"]
        photographer = photo["user"]["name"]
        img_resp = requests.get(image_url, timeout=15)
        img_resp.raise_for_status()

        os.makedirs(cache_dir, exist_ok=True)
        file_path = os.path.join(cache_dir, f"{uuid.uuid4().hex}.jpg")
        with open(file_path, "wb") as f:
            f.write(img_resp.content)
        return file_path, f"Photo by {photographer} on Unsplash"
    except Exception:
        return None


def _match_teacher_library_image(query: str, db) -> tuple[str, str | None] | None:
    """Tier 2: a teacher-uploaded image whose tags overlap the query keywords."""
    from .models import LessonImage
    query_words = set(re.findall(r"[a-z0-9]+", query.lower()))
    if not query_words:
        return None
    best_match, best_score = None, 0
    for img in db.query(LessonImage).all():
        tag_words = set(re.findall(r"[a-z0-9]+", img.tags.lower()))
        score = len(query_words & tag_words)
        if score > best_score:
            best_match, best_score = img, score
    if best_match and best_score > 0:
        return best_match.file_path, best_match.attribution_text
    return None


def resolve_slide_image(query: str, db, cache_dir: str) -> dict:
    """
    Returns {"type": "photo", "path": ..., "attribution": ...} or {"type": "generated"} -
    the renderer draws a clean icon-based graphic for the "generated" case instead of a photo.
    Tries Unsplash, then the teacher library, then falls back - never fails outright.
    """
    if not query:
        return {"type": "generated"}

    unsplash_result = _fetch_unsplash_image(query, cache_dir)
    if unsplash_result:
        path, attribution = unsplash_result
        return {"type": "photo", "path": path, "attribution": attribution}

    library_result = _match_teacher_library_image(query, db)
    if library_result:
        path, attribution = library_result
        return {"type": "photo", "path": path, "attribution": attribution}

    return {"type": "generated"}


# ---------------------------------------------------------------------------
# Rendering - python-pptx. Slide size 16:9 (13.333" x 7.5"). Each slide type gets its
# own layout, not one bullet template reused everywhere - varied per the design guidance
# (no accent-stripe clichés, proper contrast, left-aligned body text, big titles).
# ---------------------------------------------------------------------------

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)
MARGIN = Inches(0.6)

PALETTES = [
    {"accent": "0F6B62", "tint": "E6F2F0", "name": "teal"},
    {"accent": "5B3A8C", "tint": "EFE9F5", "name": "purple"},
    {"accent": "1F5C3A", "tint": "E5F0E9", "name": "forest"},
    {"accent": "B84A2E", "tint": "FBEAE5", "name": "terracotta"},
    {"accent": "7A1F3D", "tint": "F5E6EC", "name": "maroon"},
]
TEXT_COLOR = "1A1A1E"
MUTED_COLOR = "5B5B63"
BG_COLOR = "FFFFFF"
ACCENT_TEXT = "FFFFFF"


def pick_palette(seed_text: str) -> dict:
    idx = sum(ord(c) for c in seed_text) % len(PALETTES)
    return PALETTES[idx]


def _rgb(hex_str: str) -> RGBColor:
    return RGBColor.from_string(hex_str)


def _set_background(slide, color_hex: str):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = _rgb(color_hex)


def _add_textbox(slide, left, top, width, height, text, size, color, bold=False, align=PP_ALIGN.LEFT, font="Calibri"):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.color.rgb = _rgb(color)
    run.font.bold = bold
    run.font.name = font
    return box


def _add_bullet_list(slide, left, top, width, height, items, size, color, font="Calibri", spacing_after=10):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        p.space_after = Pt(spacing_after)
        run = p.add_run()
        run.text = f"•  {item}"
        run.font.size = Pt(size)
        run.font.color.rgb = _rgb(color)
        run.font.name = font
    return box


def _add_numbered_list(slide, left, top, width, height, items, size, color, font="Calibri", spacing_after=10):
    """
    Same shape as _add_bullet_list but numbered (1. 2. 3. ...) instead of bulleted -
    sequential practical/activity steps read far better numbered than as unordered
    bullets, and critically, as genuinely SEPARATE paragraphs rather than one run-on
    string with the numbers typed inline (the actual bug this was added to fix - a wall
    of "1. X. 2. Y. 3. Z." text with no real line breaks between steps).
    """
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        p.space_after = Pt(spacing_after)
        run = p.add_run()
        run.text = f"{i + 1}.  {item}"
        run.font.size = Pt(size)
        run.font.color.rgb = _rgb(color)
        run.font.name = font
    return box


def _normalize_instructions(instructions) -> list[str]:
    """
    The prompt now asks for a list of separate step strings, but this stays defensive
    against the two other shapes that could still arrive: a plain string (an old,
    already-generated deck's stored outline_json from before this change, or a rare AI
    slip that ignores the instruction) is treated as one single step rather than crashing;
    an empty/missing value returns an empty list so the caller can skip rendering cleanly.
    """
    if not instructions:
        return []
    if isinstance(instructions, list):
        return [str(step) for step in instructions if str(step).strip()]
    return [str(instructions)]


def _estimate_bullet_block_height(items: list[str], width_inches: float, font_size_pt: int, spacing_after_pt: int = 8) -> float:
    """
    Rough estimate (in inches) of how much vertical space a wrapped bullet/numbered list
    will take at a given width and font size - used to position the NEXT element
    dynamically below it, rather than at a fixed coordinate that overlaps whenever the
    list runs longer than whoever picked that fixed number expected. Not a pixel-exact
    text-layout engine (that would need to replicate pptx's own rendering), just a
    deliberately cautious estimate good enough to avoid the overlap this was built to fix.
    """
    chars_per_line = max(10, (width_inches * 72) / (0.5 * font_size_pt))
    line_height_in = (font_size_pt * 1.2) / 72
    spacing_in = spacing_after_pt / 72
    total = 0.0
    for item in items:
        lines = max(1, -(-len(item) // int(chars_per_line)))  # ceil division
        total += lines * line_height_in + spacing_in
    return total


def _pick_key_points_font_size(items: list[str], width_inches: float, available_height_inches: float) -> int:
    """
    Shrink-to-fit for a key_points list: tries the normal 15pt first, steps down only if
    the estimated block genuinely wouldn't fit the space actually available before the
    next element, down to an 11pt floor (a tight fit is still far better than an
    overlapping one).
    """
    for size in (15, 13, 11):
        if _estimate_bullet_block_height(items, width_inches, size) <= available_height_inches:
            return size
    return 11


def _content_columns(image_on_left: bool):
    """
    Text/image column positions for a content-style slide, on either side - randomized
    per slide by the caller so a deck doesn't read as "text always left, image always
    right" throughout. Returns (text_left, text_width, image_left, image_width).
    """
    text_width = Inches(6.6)
    column_gap = Inches(0.4)
    if image_on_left:
        image_left = MARGIN
        image_width = SLIDE_W - MARGIN - column_gap - text_width - MARGIN
        text_left = image_left + image_width + column_gap
    else:
        text_left = MARGIN
        image_left = text_left + text_width + column_gap
        image_width = SLIDE_W - image_left - MARGIN
    return text_left, text_width, image_left, image_width


def _add_badge(slide, left, top, text, accent_hex):
    """A small rounded-rect badge (e.g. for time allocation) - not an edge stripe, a
    self-contained shape, per the design guidance against decorative accent bars."""
    width, height = Inches(1.8), Inches(0.5)
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = _rgb(accent_hex)
    shape.line.fill.background()
    tf = shape.text_frame
    tf.word_wrap = False
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = text
    run.font.size = Pt(14)
    run.font.bold = True
    run.font.color.rgb = _rgb(ACCENT_TEXT)
    return shape


def _add_image_or_generated(slide, image_info, left, top, width, height, palette, label=""):
    """Places a real photo if resolved, otherwise a clean generated placeholder graphic
    (a tinted rounded rectangle with a simple geometric mark) - never leaves a blank gap.

    A real photo's own aspect ratio is preserved (a "contain" fit, centered within the
    given box) rather than stretched to exactly fill width x height - forcing an exact
    width AND height onto add_picture distorts any photo whose natural proportions don't
    match the box, which is what was producing the odd/unnatural-looking stretched images.
    """
    if image_info.get("type") == "photo" and os.path.exists(image_info.get("path", "")):
        try:
            fit_width, fit_height = width, height
            try:
                from PIL import Image
                with Image.open(image_info["path"]) as img:
                    img_w, img_h = img.size
                if img_w > 0 and img_h > 0:
                    box_w_in, box_h_in = width / 914400, height / 914400
                    scale = min(box_w_in / img_w, box_h_in / img_h)
                    fit_width = Inches(img_w * scale)
                    fit_height = Inches(img_h * scale)
            except Exception:
                pass  # PIL couldn't read this file's dimensions - fall back to the box size as-is (old stretch behavior) rather than lose the photo entirely

            picture_left = left + (width - fit_width) // 2
            picture_top = top + (height - fit_height) // 2
            slide.shapes.add_picture(image_info["path"], picture_left, picture_top, width=fit_width, height=fit_height)
            if image_info.get("attribution"):
                _add_textbox(
                    slide, left, top + height - Inches(0.3), width, Inches(0.3),
                    image_info["attribution"], 8, MUTED_COLOR,
                )
            return
        except Exception:
            pass  # corrupt/unreadable image file - fall through to generated placeholder

    # Generated fallback: tinted card with a simple circular mark and the topic label -
    # deliberately simple shapes only (no SVG/emoji dependency), still visually intentional.
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    card.fill.solid()
    card.fill.fore_color.rgb = _rgb(palette["tint"])
    card.line.fill.background()

    circle_size = min(width, height) * 0.35
    circle_left = left + (width - circle_size) // 2
    circle_top = top + Inches(0.25)
    circle = slide.shapes.add_shape(MSO_SHAPE.OVAL, circle_left, circle_top, circle_size, circle_size)
    circle.fill.solid()
    circle.fill.fore_color.rgb = _rgb(palette["accent"])
    circle.line.fill.background()

    if label:
        _add_textbox(
            slide, left + Inches(0.2), top + height - Inches(0.6), width - Inches(0.4), Inches(0.5),
            label, 12, palette["accent"], bold=True, align=PP_ALIGN.CENTER,
        )


def _new_slide(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])  # blank layout


def _add_title_only(slide, title_text, palette, top=Inches(0.5)):
    _add_textbox(slide, MARGIN, top, SLIDE_W - 2 * MARGIN, Inches(0.9), title_text, 32, palette["accent"], bold=True)


def _build_title_slide(prs, data, palette):
    slide = _new_slide(prs)
    _set_background(slide, palette["accent"])
    _add_textbox(slide, MARGIN, Inches(2.6), SLIDE_W - 2 * MARGIN, Inches(1.5),
                 data.get("title", "Lesson"), 44, ACCENT_TEXT, bold=True, align=PP_ALIGN.LEFT)
    if data.get("subtitle"):
        _add_textbox(slide, MARGIN, Inches(4.0), SLIDE_W - 2 * MARGIN, Inches(0.8),
                     data["subtitle"], 20, ACCENT_TEXT, align=PP_ALIGN.LEFT)


def _build_retrieval_slide(prs, data, palette, heading_default):
    slide = _new_slide(prs)
    _set_background(slide, BG_COLOR)
    _add_title_only(slide, data.get("title", heading_default), palette)
    if data.get("time_minutes"):
        _add_badge(slide, SLIDE_W - MARGIN - Inches(1.8), Inches(0.55), f"{data['time_minutes']} min", palette["accent"])
    _add_bullet_list(slide, MARGIN, Inches(1.8), SLIDE_W - 2 * MARGIN, Inches(4.8),
                      data.get("questions", []), 22, TEXT_COLOR, spacing_after=16)


def _build_objectives_slide(prs, data, palette, review=False):
    slide = _new_slide(prs)
    _set_background(slide, BG_COLOR)
    heading = "Have We Met Today's Objectives?" if review else "Learning Objectives"
    _add_title_only(slide, heading, palette)

    if review:
        items = data.get("objectives_check", [])
        top = Inches(1.8)
        for item in items:
            _add_textbox(slide, MARGIN, top, SLIDE_W - 2 * MARGIN, Inches(0.5),
                         f"✓  {item.get('objective', '')}", 20, TEXT_COLOR, bold=True)
            top += Inches(0.55)
            _add_textbox(slide, MARGIN + Inches(0.3), top, SLIDE_W - 2 * MARGIN - Inches(0.3), Inches(0.5),
                         item.get("self_assess_prompt", ""), 16, MUTED_COLOR)
            top += Inches(0.85)
    else:
        _add_bullet_list(slide, MARGIN, Inches(1.8), SLIDE_W - 2 * MARGIN, Inches(4.8),
                          data.get("objectives", []), 22, TEXT_COLOR, spacing_after=18)


def _build_content_slide(prs, data, palette, image_info, image_on_left=False):
    slide = _new_slide(prs)
    _set_background(slide, BG_COLOR)
    _add_title_only(slide, data.get("title", "Content"), palette)

    text_left, text_width, image_left, image_width = _content_columns(image_on_left)
    image_height = Inches(4.6)

    if data.get("guiding_question"):
        _add_textbox(slide, text_left, Inches(1.7), text_width, Inches(0.9),
                     data["guiding_question"], 19, palette["accent"], bold=True)

    key_points = data.get("key_points", [])
    key_points_top_in = 2.75
    if key_points:
        # Available room before check_question needs to start, with a small gap and space
        # for check_question's own text before the slide's bottom margin.
        available_in = (7.5 - 0.6) - key_points_top_in - 0.9
        font_size = _pick_key_points_font_size(key_points, text_width / 914400, available_in)
        _add_bullet_list(slide, text_left, Inches(key_points_top_in), text_width, Inches(available_in),
                          key_points, font_size, TEXT_COLOR, spacing_after=8)
        estimated_height = _estimate_bullet_block_height(key_points, text_width / 914400, font_size)
        check_question_top = min(key_points_top_in + estimated_height + 0.2, 7.5 - 0.6 - 0.9)
    else:
        check_question_top = key_points_top_in

    if data.get("check_question"):
        _add_textbox(slide, text_left, Inches(check_question_top), text_width, Inches(0.9),
                     f"Check: {data['check_question']}", 14, MUTED_COLOR)

    _add_image_or_generated(slide, image_info, image_left, Inches(1.7), image_width, image_height, palette,
                             label=data.get("image_query", ""))


def _build_content_question_slide(prs, data, palette, image_info, image_on_left=False):
    """
    First half of a two-part content slide: title, guiding question, and the image - no
    key points or check question yet. Built so a teacher can present this, pose the
    question, and hear from students BEFORE advancing to the reveal slide.

    This achieves "don't show the answer until I click" via a genuine next-slide advance,
    not an in-slide fade/appear animation. python-pptx has no supported API for native
    PowerPoint build animations - they'd have to be hand-written as raw OOXML timing XML,
    which is undocumented, easy to get subtly wrong, and can't be verified here (no real
    PowerPoint install to open the result in and confirm it isn't corrupted). Two ordinary
    slides is something python-pptx handles completely reliably, and the effect during an
    actual presentation - press to advance, the answer appears - is the same either way.
    """
    slide = _new_slide(prs)
    _set_background(slide, BG_COLOR)
    _add_title_only(slide, data.get("title", "Content"), palette)

    text_left, text_width, image_left, image_width = _content_columns(image_on_left)
    image_height = Inches(4.6)

    if data.get("guiding_question"):
        _add_textbox(slide, text_left, Inches(2.6), text_width, Inches(2.2),
                     data["guiding_question"], 26, palette["accent"], bold=True)

    _add_image_or_generated(slide, image_info, image_left, Inches(1.7), image_width, image_height, palette,
                             label=data.get("image_query", ""))


def _build_content_answer_slide(prs, data, palette, image_info, image_on_left=False):
    """
    Second half - the reveal, advanced to once the teacher has heard student answers.
    Repeats the guiding question smaller at the top for continuity, then the key points
    and check question - the same content the original single-slide version showed, just
    moved to this second slide.
    """
    slide = _new_slide(prs)
    _set_background(slide, BG_COLOR)
    _add_title_only(slide, data.get("title", "Content"), palette)

    text_left, text_width, image_left, image_width = _content_columns(image_on_left)
    image_height = Inches(4.6)

    if data.get("guiding_question"):
        _add_textbox(slide, text_left, Inches(1.7), text_width, Inches(0.7),
                     data["guiding_question"], 16, MUTED_COLOR, bold=True)

    key_points = data.get("key_points", [])
    key_points_top_in = 2.5
    if key_points:
        available_in = (7.5 - 0.6) - key_points_top_in - 0.9
        font_size = _pick_key_points_font_size(key_points, text_width / 914400, available_in)
        _add_bullet_list(slide, text_left, Inches(key_points_top_in), text_width, Inches(available_in),
                          key_points, font_size, TEXT_COLOR, spacing_after=8)
        estimated_height = _estimate_bullet_block_height(key_points, text_width / 914400, font_size)
        check_question_top = min(key_points_top_in + estimated_height + 0.2, 7.5 - 0.6 - 0.9)
    else:
        check_question_top = key_points_top_in

    if data.get("check_question"):
        _add_textbox(slide, text_left, Inches(check_question_top), text_width, Inches(0.9),
                     f"Check: {data['check_question']}", 14, MUTED_COLOR)

    _add_image_or_generated(slide, image_info, image_left, Inches(1.7), image_width, image_height, palette,
                             label=data.get("image_query", ""))


def _build_real_life_slide(prs, data, palette, image_info):
    slide = _new_slide(prs)
    _set_background(slide, BG_COLOR)
    _add_title_only(slide, data.get("title", "Real-Life Example"), palette)

    image_height = Inches(3.4)
    _add_image_or_generated(slide, image_info, MARGIN, Inches(1.6), SLIDE_W - 2 * MARGIN, image_height,
                             palette, label=data.get("image_query", ""))

    top = Inches(1.6) + image_height + Inches(0.3)
    if data.get("scenario"):
        _add_textbox(slide, MARGIN, top, SLIDE_W - 2 * MARGIN, Inches(1.0), data["scenario"], 16, TEXT_COLOR)
        top += Inches(1.0)
    if data.get("discussion_question"):
        _add_textbox(slide, MARGIN, top, SLIDE_W - 2 * MARGIN, Inches(0.6),
                     f"Discuss: {data['discussion_question']}", 16, palette["accent"], bold=True)


def _build_activity_slide(prs, data, palette, kind="activity"):
    """Activities/practicals get a full tinted background (not an edge stripe - the design
    guidance explicitly warns against accent stripes as an AI-slide cliché) so they read as
    visually distinct at a glance when flipping through the deck."""
    slide = _new_slide(prs)
    _set_background(slide, palette["tint"])

    label = "PRACTICAL" if kind == "practical" else "ACTIVITY"
    _add_textbox(slide, MARGIN, Inches(0.5), Inches(4), Inches(0.4), label, 14, palette["accent"], bold=True)
    _add_textbox(slide, MARGIN, Inches(0.95), SLIDE_W - 2 * MARGIN, Inches(0.9),
                 data.get("title", ""), 28, TEXT_COLOR, bold=True)

    if data.get("time_minutes"):
        _add_badge(slide, SLIDE_W - MARGIN - Inches(1.8), Inches(0.55), f"{data['time_minutes']} min", palette["accent"])

    if kind == "activity" and data.get("grouping"):
        _add_textbox(slide, MARGIN, Inches(1.75), Inches(4), Inches(0.4),
                     f"Grouping: {data['grouping']}", 14, MUTED_COLOR)

    top = Inches(2.3)
    instruction_steps = _normalize_instructions(data.get("instructions"))
    if instruction_steps:
        _add_numbered_list(slide, MARGIN, top, SLIDE_W - 2 * MARGIN, Inches(3.4), instruction_steps, 16, TEXT_COLOR, spacing_after=8)
        top += Inches(1.5)

    if kind == "practical" and data.get("materials_needed"):
        _add_textbox(slide, MARGIN, Inches(6.2), SLIDE_W - 2 * MARGIN, Inches(0.8),
                     f"Materials: {data['materials_needed']}", 14, MUTED_COLOR)


def render_lesson_pptx(outline: dict, db, image_cache_dir: str, output_path: str) -> list[str]:
    """
    Builds the actual .pptx file from the AI-generated outline. Returns a list of any
    slide titles where image resolution fell through to the generated placeholder (so the
    caller can report "N of M slides used real images" rather than claiming full success
    silently when photos weren't available).
    """
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H

    palette = pick_palette(outline.get("lesson_title", "lesson"))
    generated_fallback_slides = []

    for slide_data in outline.get("slides", []):
        slide_type = slide_data.get("type")

        if slide_type == "title":
            _build_title_slide(prs, slide_data, palette)

        elif slide_type == "retrieval_starter":
            _build_retrieval_slide(prs, slide_data, palette, "Do Now: Last Lesson Recap")

        elif slide_type == "objectives":
            _build_objectives_slide(prs, slide_data, palette, review=False)

        elif slide_type == "content":
            image_info = slide_data.get("resolved_image") or resolve_slide_image(slide_data.get("image_query", ""), db, image_cache_dir)
            if image_info["type"] == "generated":
                generated_fallback_slides.append(slide_data.get("title", "Content slide"))
            # Decided once per content entry (not per builder call) so a question/reveal
            # pair keeps the image on the same side across both slides - varies ACROSS
            # different content entries in the deck instead of every slide defaulting to
            # the same side, which was the actual complaint.
            image_on_left = random.choice([True, False])
            if slide_data.get("guiding_question"):
                # Two slides so the teacher can pose the question and hear from students
                # before advancing to reveal the key points - see the two builder
                # functions' own docstrings for why this is a next-slide advance rather
                # than an in-slide animation.
                _build_content_question_slide(prs, slide_data, palette, image_info, image_on_left)
                _build_content_answer_slide(prs, slide_data, palette, image_info, image_on_left)
            else:
                # No guiding question was generated for this slide (shouldn't normally
                # happen per the prompt, but handled safely) - nothing to reveal, so this
                # stays a single slide exactly as before.
                _build_content_slide(prs, slide_data, palette, image_info, image_on_left)

        elif slide_type == "real_life_example":
            image_info = slide_data.get("resolved_image") or resolve_slide_image(slide_data.get("image_query", ""), db, image_cache_dir)
            if image_info["type"] == "generated":
                generated_fallback_slides.append(slide_data.get("title", "Real-life example"))
            _build_real_life_slide(prs, slide_data, palette, image_info)

        elif slide_type == "activity":
            _build_activity_slide(prs, slide_data, palette, kind="activity")

        elif slide_type == "practical":
            _build_activity_slide(prs, slide_data, palette, kind="practical")

        elif slide_type == "retrieval_plenary":
            _build_retrieval_slide(prs, slide_data, palette, "Check Your Understanding")

        elif slide_type == "objectives_review":
            _build_objectives_slide(prs, slide_data, palette, review=True)

    prs.save(output_path)
    return generated_fallback_slides
