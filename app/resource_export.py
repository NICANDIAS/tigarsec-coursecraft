"""
Renders the printable resources a lesson's activities/practicals actually need (sample
data tables, cards to cut up, scenario briefs, worksheets) into a downloadable Word
document - generated fresh from the deck's already-stored outline_json each time it's
requested, so it stays available (and re-downloadable) any time after the PowerPoint
itself was generated, without needing to regenerate anything.
"""
from docx import Document


def extract_resources(outline_json: dict) -> list[dict]:
    """Pulls out every slide that actually has a resource, in slide order."""
    resources = []
    for slide in outline_json.get("slides", []):
        resource = slide.get("resource")
        if resource and resource.get("needed"):
            resources.append({
                "slide_title": slide.get("title", ""),
                "resource_title": resource.get("title", "Resource"),
                "content": resource.get("content", ""),
            })
    return resources


def differentiated_activities(outline_json: dict) -> list[dict]:
    """Stage 10: activities that have support/stretch versions, in slide order."""
    out = []
    for slide in (outline_json or {}).get("slides", []):
        versions = slide.get("differentiation")
        if slide.get("type") in ("activity", "practical") and versions:
            out.append({"title": slide.get("title", "Activity"), "core": slide.get("instructions"), **versions})
    return out


def worksheet_count(outline_json: dict) -> int:
    return len(extract_resources(outline_json or {})) + len(differentiated_activities(outline_json))


def _steps(doc, steps):
    if isinstance(steps, str):
        steps = [steps]
    for n, step in enumerate(steps or [], start=1):
        doc.add_paragraph(f"{n}. {step}")


def render_resources_docx(outline_json: dict, lesson_title: str, output_path: str) -> int:
    """Returns the number of resources written, so the caller can decide whether there's
    anything worth offering as a download at all. Stage 10: also writes each activity's
    Support / Core / Stretch worksheets, when the deck has them."""
    resources = extract_resources(outline_json)
    differentiated = differentiated_activities(outline_json)
    if not resources and not differentiated:
        return 0

    doc = Document()
    doc.add_heading(f"Activity Resources: {lesson_title}", level=1)
    doc.add_paragraph(
        "Print and hand out as needed for the activities/practicals in this lesson - "
        "generated to match the PowerPoint exactly."
    )

    for i, res in enumerate(resources):
        if i > 0:
            doc.add_page_break()
        doc.add_heading(res["resource_title"], level=2)
        italic_para = doc.add_paragraph()
        italic_run = italic_para.add_run(f"For: {res['slide_title']}")
        italic_run.italic = True
        for line in res["content"].split("\n"):
            doc.add_paragraph(line)

    for act in differentiated:
        for label, version in (("Support", act.get("support")), ("Core", {"steps": act.get("core")}), ("Stretch", act.get("stretch"))):
            if not version or not version.get("steps"):
                continue
            if resources or act is not differentiated[0] or label != "Support":
                doc.add_page_break()
            doc.add_heading(f"{act['title']} - {label}", level=2)
            note = doc.add_paragraph()
            note.add_run({"Support": "Step-by-step version with help to get started.",
                          "Core": "The task as set on the slide.",
                          "Stretch": "Extension - for when you have finished, or are aiming higher."}[label]).italic = True
            doc.add_paragraph("Name: ____________________________")
            _steps(doc, version.get("steps"))
            if label == "Support" and version.get("scaffolds"):
                doc.add_heading("Help", level=3)
                for s in version["scaffolds"]:
                    doc.add_paragraph(s, style="List Bullet")
            if label == "Stretch" and version.get("challenge"):
                doc.add_heading("Challenge", level=3)
                doc.add_paragraph(version["challenge"])

    doc.save(output_path)
    return len(resources) + len(differentiated)
