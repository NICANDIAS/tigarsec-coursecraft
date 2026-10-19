"""
Stage 12 (1): the class tracking sheet - every student against every criterion in every unit,
filled in automatically from their marked work (no more hand-kept Excel trackers). Pure
functions, no database: the router gathers the marks, this lays them out.

A cell is the student's result on that criterion in their LATEST marked work for the unit:
"met", "partly", "not" or "" (not assessed yet). Each unit also shows the unit grade (P/M/D
specs) or how many criteria are met. "At risk" flags students a lecturer should look at now.
"""
import io

from .coverage import classify_mark
from .marking_insights import unit_criteria, match_criterion

STATE = {"met": "met", "partly_met": "partly", "not_met": "not"}
SYMBOL = {"met": "✓", "partly": "½", "not": "✗", "": ""}
GRADE_RANK = {"Not yet achieved": 0, "Pass": 1, "Merit": 2, "Distinction": 3}


def build_tracker(units: list, students: list, marks_by_unit: dict) -> dict:
    """units: the spec's units. students: [{"id", "name"}] (active, in class order).
    marks_by_unit: {unit name: {student id: {"submission_id", "grade", "marks": [{"criterion_ref", "final_mark"}]}}}
    -> {"units": [...], "rows": [...], "unit_summary": {...}} - only units that have marked work."""
    shown = []
    for i, u in enumerate(units or []):
        name = u.get("name")
        if marks_by_unit.get(name):
            shown.append({"index": i, "name": name, "criteria": [c["code"] for c in unit_criteria(u)], "_crit": unit_criteria(u)})
    rows = []
    for st in students:
        cells, risk = {}, []
        for u in shown:
            work = marks_by_unit[u["name"]].get(st["id"])
            crit = {code: "" for code in u["criteria"]}
            if work:
                for m in work["marks"]:
                    c = match_criterion(m.get("criterion_ref", ""), u["_crit"])
                    if c:
                        crit[c["code"]] = STATE.get(classify_mark(m.get("final_mark")), crit[c["code"]])
            met = sum(1 for v in crit.values() if v == "met")
            assessed = sum(1 for v in crit.values() if v)
            cells[u["name"]] = {"crit": crit, "grade": (work or {}).get("grade"), "submission_id": (work or {}).get("submission_id"),
                                "met": met, "assessed": assessed, "has_work": bool(work)}
            if work and (work.get("grade") == "Not yet achieved" or (assessed and met == 0)):
                risk.append(f"not yet passing {u['name']}")
        rows.append({"student": st, "cells": cells, "risk": risk})
    # missing work: a unit most of the class has handed in, but this student hasn't
    class_size = max(len(students), 1)
    for u in shown:
        handed_in = sum(1 for r in rows if r["cells"][u["name"]]["has_work"])
        if handed_in * 2 >= class_size:
            for r in rows:
                if not r["cells"][u["name"]]["has_work"]:
                    r["risk"].append(f"no marked work for {u['name']}")
    summary = {}
    for u in shown:
        grades = [r["cells"][u["name"]]["grade"] for r in rows if r["cells"][u["name"]]["has_work"]]
        summary[u["name"]] = {
            "handed_in": len(grades),
            "grades": {g: grades.count(g) for g in GRADE_RANK if grades.count(g)},
            "criteria_met_pct": {code: _pct(sum(1 for r in rows if r["cells"][u["name"]]["crit"][code] == "met"),
                                            sum(1 for r in rows if r["cells"][u["name"]]["crit"][code]))
                                 for code in u["criteria"]},
        }
    for u in shown:
        u.pop("_crit", None)
    return {"units": shown, "rows": rows, "unit_summary": summary,
            "at_risk": sum(1 for r in rows if r["risk"])}


def _pct(n, d):
    return round(100 * n / d) if d else None


def tracker_xlsx(course_name: str, tracker: dict) -> bytes:
    """The tracking sheet as an Excel file, coloured like the page - for a head of department
    or a team that still wants a spreadsheet."""
    from openpyxl import Workbook
    from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    fills = {"met": PatternFill("solid", fgColor="C6EFCE"), "partly": PatternFill("solid", fgColor="FFEB9C"),
             "not": PatternFill("solid", fgColor="FFC7CE")}
    thin = Side(style="thin", color="BBBBBB")
    wb = Workbook()
    ws = wb.active
    ws.title = "Tracker"
    ws["A1"] = f"{course_name} - tracking sheet"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = "✓ met   ½ partly met   ✗ not met   (blank = not assessed yet). From each student's latest marked work."
    # header rows: unit names over their criteria, then the criterion codes
    col = 2
    ws.cell(row=4, column=1, value="Student").font = Font(bold=True)
    grade_cols = []
    for u in tracker["units"]:
        span = len(u["criteria"]) + 1
        ws.cell(row=3, column=col, value=u["name"]).font = Font(bold=True)
        if span > 1:
            ws.merge_cells(start_row=3, start_column=col, end_row=3, end_column=col + span - 1)
        for j, code in enumerate(u["criteria"]):
            c = ws.cell(row=4, column=col + j, value=code)
            c.font = Font(bold=True)
            c.alignment = Alignment(horizontal="center")
        ws.cell(row=4, column=col + span - 1, value="Grade").font = Font(bold=True)
        grade_cols.append(col + span - 1)
        col += span
    ws.cell(row=4, column=col, value="At risk").font = Font(bold=True)
    for r_i, row in enumerate(tracker["rows"], start=5):
        ws.cell(row=r_i, column=1, value=row["student"]["name"])
        col = 2
        for u in tracker["units"]:
            cell = row["cells"][u["name"]]
            for j, code in enumerate(u["criteria"]):
                state = cell["crit"][code]
                c = ws.cell(row=r_i, column=col + j, value=SYMBOL[state])
                c.alignment = Alignment(horizontal="center")
                c.border = Border(left=thin, right=thin, top=thin, bottom=thin)
                if state:
                    c.fill = fills[state]
            grade = cell["grade"] or (f"{cell['met']}/{len(u['criteria'])} met" if cell["has_work"] else "")
            ws.cell(row=r_i, column=col + len(u["criteria"]), value=grade).font = Font(bold=True)
            col += len(u["criteria"]) + 1
        ws.cell(row=r_i, column=col, value="; ".join(row["risk"]))
    ws.column_dimensions["A"].width = 24
    for i in range(2, ws.max_column + 1):
        ws.column_dimensions[get_column_letter(i)].width = 6
    for g in grade_cols:
        ws.column_dimensions[get_column_letter(g)].width = 14
    ws.column_dimensions[get_column_letter(ws.max_column)].width = 40
    ws.freeze_panes = "B5"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
