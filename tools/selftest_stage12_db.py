"""
Stage 12 self-test: the tracking sheet (page + Excel download) and personal revision after a
live quiz - through the REAL page functions, with a THROWAWAY database in /tmp. Your real
database is not opened. No AI is used.

Run (from ~/projects/coursecraft):
  docker compose run --rm --no-deps -v "$PWD":/app --entrypoint python app tools/selftest_stage12_db.py
Expected last line: "ALL STAGE 12 CHECKS PASSED (14)"
"""
import io
import os
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
os.chdir(ROOT)

TEMP_DB = "/tmp/coursecraft_stage12_selftest.db"
if os.path.exists(TEMP_DB):
    os.remove(TEMP_DB)
os.environ["DATABASE_URL"] = f"sqlite:///{TEMP_DB}"

from starlette.requests import Request  # noqa: E402
from app import database  # noqa: E402

assert str(database.engine.url).endswith("coursecraft_stage12_selftest.db"), "SAFETY STOP: not the throwaway database"

from app.database import Base, SessionLocal, engine  # noqa: E402
from app.models import (Teacher, Course, CourseTeacher, SpecVersion, Student, Submission, MarkingRun, CriterionMark,  # noqa: E402
                        QuizSet, LiveSession, LiveParticipant)
from app.routers import tracker as tk, live as lv, courses as cr  # noqa: E402

RESULTS = []


def check(name, condition, detail=""):
    RESULTS.append(bool(condition))
    print(f"{'PASS' if condition else 'FAIL'}  {name}" + (f"  ({detail})" if detail and not condition else ""))


def request(path, cookie=None, query=""):
    headers = [(b"cookie", cookie.encode())] if cookie else []
    return Request({"type": "http", "method": "GET", "path": path, "query_string": query.encode(), "headers": headers,
                    "server": ("test", 80), "scheme": "http", "root_path": ""})


UNIT = {"name": "Unit 2 Networking", "code": "2", "learning_outcomes": [{"code": "A", "title": "Network types", "criteria": [
    {"code": "P1", "description": "Explain two network topologies."},
    {"code": "M1", "description": "Compare the performance of two topologies."},
    {"code": "D1", "description": "Evaluate the choice of topology for a given business."}]}]}

Base.metadata.create_all(bind=engine)
db = SessionLocal()
lecturer = Teacher(name="Self Test", email="selftest@example.invalid", hashed_password="x", role="lecturer")
other = Teacher(name="Other", email="other@example.invalid", hashed_password="x", role="lecturer")
db.add_all([lecturer, other])
db.commit()
spec = SpecVersion(uploaded_by_teacher_id=lecturer.id, label="BTEC IT", raw_text="x", parse_status="parsed", structured_json={"units": [UNIT]})
db.add(spec)
db.commit()
course = Course(teacher_id=lecturer.id, name="BTEC IT Y1", spec_version_id=spec.id)
db.add(course)
db.commit()
db.add(CourseTeacher(course_id=course.id, teacher_id=lecturer.id))
db.commit()

# ------------------------------------------------------------------ the tracking sheet
empty = tk.tracker_page(course.id, request("/t"), "all", 0, db, lecturer).body.decode()
check("before any marking, the page says nothing is marked yet", "Nothing marked yet on this course" in empty)

MARKS = {"Amira Khan": ("Merit", [("P1", "Met"), ("M1", "Met"), ("D1", "Not met")]),
         "Ben Okafor": ("Not yet achieved", [("P1", "Partly met"), ("M1", "Not met"), ("D1", "Not met")]),
         "Chloe Price": ("Pass", [("P1", "Met"), ("M1", "Not met"), ("D1", "Not met")]),
         "Dan Fry": None}
subs = {}
for name, marked in MARKS.items():
    st = Student(course_id=course.id, name=name)
    db.add(st)
    db.commit()
    if not marked:
        continue
    sub = Submission(student_id=st.id, spec_version_id=spec.id, target_unit_name=UNIT["name"], original_filename="work.docx",
                     status="needs_review", achieved_grade=marked[0])
    db.add(sub)
    db.commit()
    run = MarkingRun(submission_id=sub.id, status="needs_review")
    db.add(run)
    db.commit()
    db.add_all([CriterionMark(marking_run_id=run.id, criterion_ref=ref, ai_mark=m, final_mark=m, final_comment="") for ref, m in marked[1]])
    db.commit()
    subs[name] = sub.id

page = tk.tracker_page(course.id, request("/t"), "all", 0, db, lecturer).body.decode()
check("every student is on the sheet, with their grade", all(n in page for n in MARKS) and "Merit" in page and "Not yet achieved" in page)
check("each result links to that piece of work", f'/submissions/{subs["Amira Khan"]}/review' in page)
check("results are coloured met / partly / not met", 'class="c-met"' in page and 'class="c-partly"' in page and 'class="c-not"' in page)
check("the class % met row is worked out (P1: 2 of 3)", "67%" in page)
check("students at risk are flagged with the reason",
      "not yet passing Unit 2 Networking" in page and "no marked work for Unit 2 Networking" in page)
risk_page = tk.tracker_page(course.id, request("/t"), "all", 1, db, lecturer).body.decode()
check("'Only students at risk' hides the others", "Ben Okafor" in risk_page and "Dan Fry" in risk_page
      and "/students/" + str(db.query(Student).filter(Student.name == "Amira Khan").first().id) + '"' not in risk_page)
check("another lecturer cannot open it", tk.tracker_page(course.id, request("/t"), "all", 0, db, other).status_code == 303)

download = tk.tracker_download(course.id, db, lecturer)
from openpyxl import load_workbook  # noqa: E402
ws = load_workbook(io.BytesIO(download.body)).active
check("the Excel download has the same sheet", download.media_type.endswith("spreadsheetml.sheet")
      and ws["A5"].value == "Amira Khan" and ws["B5"].value == "✓" and ws["E5"].value == "Merit")
course_page = cr.course_detail(course.id, request(f"/courses/{course.id}"), "", db, lecturer).body.decode()
check("the course page links to the tracking sheet", f"/courses/{course.id}/tracker" in course_page)

# ------------------------------------------------------------------ personal revision after a live quiz
quiz = QuizSet(course_id=course.id, teacher_id=lecturer.id, title="Networks quiz", seconds_per_question=20, status="ready",
               questions_json=[{"question": "HTTPS port?", "choices": ["80", "443"], "answer": 1, "explanation": "443 is HTTPS."},
                               {"question": "A star has a central...?", "choices": ["switch", "ring"], "answer": 0, "explanation": ""}])
db.add(quiz)
db.commit()
lv.live_start_from_quiz(quiz.id, db, lecturer)
db.expire_all()
session = db.query(LiveSession).filter(LiveSession.quiz_set_id == quiz.id).first()
lv.live_begin(session.id, db, lecturer)
cookie = lv.join_submit(request("/join"), session.join_code, "Amira", db).headers["set-cookie"].split(";")[0]
lv.live_set_screen(session.id, 1, db, lecturer)
db.expire_all()
key = db.get(LiveSession, session.id).screens_json[1]["questions"][0]["key"]
answered = lv.student_answer(session.join_code, request("/a", cookie), 1, key, "", 0, db)  # wrong; Q2 left unanswered
during = lv.student_revision(session.id, request("/r", cookie), db)
check("no revision while the quiz is still running (no answers given away)", during["questions"] == [] and answered == {"ok": True}, answered)
lv.live_end(session.id, db, lecturer)
after = lv.student_revision(session.id, request("/r", cookie), db)
check("after the quiz, the student gets the questions they got wrong or missed",
      [q["question"] for q in after["questions"]] == ["HTTPS port?", "A star has a central...?"] and after["total"] == 2)
check("...with the right answer and the explanation", after["questions"][0]["answer"] == 1 and after["questions"][0]["explanation"] == "443 is HTTPS.")
nobody = lv.student_revision(session.id, request("/r"), db)
check("someone who didn't join gets nothing", getattr(nobody, "status_code", 200) == 404)

db.close()
os.remove(TEMP_DB)
print()
if all(RESULTS):
    print(f"ALL STAGE 12 CHECKS PASSED ({len(RESULTS)})")
else:
    print(f"{RESULTS.count(False)} OF {len(RESULTS)} CHECKS FAILED - send me this output")
    sys.exit(1)
