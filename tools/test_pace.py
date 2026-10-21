"""
Tests for teaching at the class's pace: a lesson carried on into the next week, everything after it
moving along the college calendar, combining two later weeks, and the printed scheme showing it all.
Builds a small college-style template (a Word file) in /tmp - no real files are touched.

Run (from ~/projects/coursecraft):
  docker compose run --rm --no-deps -v "$PWD":/app --entrypoint python app tools/test_pace.py
Expected last line: "ALL 9 TESTS PASSED"
"""
import datetime
import os
import sys
import tempfile
from pathlib import Path
from types import SimpleNamespace

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
from docx import Document  # noqa: E402
from app import pace as pc  # noqa: E402
from app.sow_calendar import scheme_calendar, scheme_now  # noqa: E402
from app.sow_docx_fill import fill_sow_template_docx  # noqa: E402
from app.term_plan import timeline_rows, quick_pick  # noqa: E402
from app.live_activities import start_screen_for_slide  # noqa: E402

TESTS = []
D = lambda m, d: datetime.date(2026, m, d)  # noqa: E731
TMP = tempfile.mkdtemp(prefix="cc_pace_")


def test(fn):
    TESTS.append(fn)
    return fn


CALENDAR = [("8-9-2026", "Induction Week"), ("15-9-2026", ""), ("22-9-2026", ""), ("29-9-2026", ""), ("6-10-2026", ""),
            ("13-10-2026", ""), ("19-10-2026", "Staff Conference Day (Friday)"), ("26-10-2026", "HALF TERM"),
            ("2-11-2026", ""), ("9-11-2026", ""), ("16-11-2026", "")]


def make_template():
    doc = Document()
    table = doc.add_table(rows=1, cols=4)
    for cell, text in zip(table.rows[0].cells, ["Week commencing", "Topic", "Content", "Assessment"]):
        cell.text = text
    for date, label in CALENDAR:
        row = table.add_row().cells
        row[0].text, row[1].text = date, label
    path = os.path.join(TMP, "college.docx")
    doc.save(path)
    return path


PATH = make_template()   # 9 teaching rows: 15 Sep ... 19 Oct, then half term, then 2, 9, 16 Nov
WEEKS_JSON = [{"week": n, "title": t, "learning_outcomes": lo} for n, t, lo in [
    (1, "Threats", ["1.1"]), (2, "Social engineering", ["1.2"]), (3, "Threat actors", ["1.3"]), (4, "Risk stages", ["2.1"]),
    (5, "CVE and CVSS", ["2.2"]), (6, "Mitigations", ["2.3"]), (7, "Risk types", ["2.4", "2.5"]), (8, "Risk measures", ["2.5"]),
    (9, "Assessment week", ["2.1", "2.2"])]]


def sow(pace=None, anchor=None):
    return SimpleNamespace(format_sample_file_path=PATH, append_eligible_labels=None, weeks_json=WEEKS_JSON,
                           start_date=None, end_date=None, week_anchor_json=anchor, pace_json=pace, planned_assessments=None)


CARRIED = {"extensions": [{"week": 5, "deck_id": 12, "from_slide": 14, "reason": "CVSS vectors", "at": "2026-10-20"}], "merges": []}


@test
def without_changes_the_plan_follows_the_template():
    cal = scheme_calendar(sow())
    assert cal["week_dates"][5] == D(10, 13) and cal["week_dates"][6] == D(10, 19) and cal["week_dates"][7] == D(11, 2)
    assert cal["overflow"] == 0


@test
def carrying_on_week_5_moves_everything_after_it_one_teaching_week():
    cal = scheme_calendar(sow(CARRIED))
    parts = [(t["week"], t["part"], t["date"]) for t in cal["timeline"] if t["kind"] == "week"]
    assert (5, 1, D(10, 13)) in parts and (5, 2, D(10, 19)) in parts
    assert cal["week_dates"][6] == D(11, 2) and cal["week_dates"][7] == D(11, 9)      # half term still skipped
    assert cal["overflow"] == 1                                                          # 10 weeks of teaching, 9 rows


@test
def this_week_knows_it_is_week_5_continued():
    now, _ = scheme_now(sow(CARRIED), None, D(10, 21))
    assert now["state"] == "teaching" and now["week"] == 5 and now["part"] == 2
    assert pc.resume_slide(CARRIED, 5, 2) == 14 and pc.resume_slide(CARRIED, 5, 1) is None


@test
def the_planning_page_shows_the_continued_week_and_never_ticks_it():
    now, cal = scheme_now(sow(CARRIED), None, D(10, 14))
    rows = timeline_rows(cal["timeline"], WEEKS_JSON, now, {1, 2, 3, 4, 5}, D(10, 14), CARRIED)
    cont = next(r for r in rows if r.get("continued"))
    assert cont["week"] == 5 and cont["title"] == "CVE and CVSS (continued)" and cont["resume"] == 14 and cont["made"]
    assert quick_pick(rows, now["week"], "rest") == [6, 7, 8, 9]


@test
def combining_two_later_weeks_wins_the_week_back():
    ideas = pc.suggest_merges(WEEKS_JSON, 5, CARRIED)
    assert ideas[0]["a"] == 7 and ideas[0]["b"] == 8 and ideas[0]["shared"] == ["2.5"]   # they share a criterion
    assert all(9 not in (i["a"], i["b"]) for i in ideas)                                 # never the assessment week
    merged = dict(CARRIED, merges=[[7, 8]])
    cal = scheme_calendar(sow(merged))
    assert cal["overflow"] == 0 and 8 not in cal["week_dates"]
    seven = next(t for t in cal["timeline"] if t["kind"] == "week" and t["week"] == 7)
    assert seven["with"] == [8]


@test
def the_printed_scheme_shows_how_it_was_really_taught():
    printed = pc.paced_weeks(WEEKS_JSON, dict(CARRIED, merges=[[7, 8]]))
    assert [w["title"] for w in printed] == ["Threats", "Social engineering", "Threat actors", "Risk stages", "CVE and CVSS",
                                             "CVE and CVSS (continued)", "Mitigations", "Risk types + Risk measures", "Assessment week"]
    assert printed[5]["activities"][0].endswith("from slide 14") and printed[7]["learning_outcomes"] == ["2.4", "2.5"]
    out = os.path.join(TMP, "printed.docx")
    fill_sow_template_docx(PATH, printed, {}, out)
    rows = Document(out).tables[0].rows
    assert "CVE and CVSS (continued)" in rows[7].cells[1].text and "Staff Conference Day" in rows[7].cells[1].text
    assert "HALF TERM" in rows[8].cells[1].text and "Mitigations" in rows[9].cells[1].text


@test
def without_a_template_weeks_follow_the_start_date_and_carry_on_too():
    plain = SimpleNamespace(format_sample_file_path=None, weeks_json=WEEKS_JSON, start_date=D(9, 14), end_date=D(11, 9),
                            week_anchor_json=None, pace_json=CARRIED, append_eligible_labels=None)
    cal = scheme_calendar(plain)
    assert cal["week_dates"][6] == D(10, 26) and cal["overflow"] == 1   # 10 weeks from 14 Sep run to 16 Nov; it ends 9 Nov
    now, _ = scheme_now(plain, None, D(10, 21))
    assert now["week"] == 5 and now["part"] == 2 and now["source"] == "start_date"


@test
def the_class_signals_and_the_suggestion():
    screens = [{"kind": "objectives", "questions": [{"key": "c1", "kind": "confidence", "prompt": "I can read a CVSS vector",
                                                     "choices": ["Not yet", "Getting there", "Got it"]}]},
               {"kind": "quiz_question", "questions": [{"key": "q1", "kind": "mcq", "prompt": "Base score?", "choices": ["a", "b"], "answer": 0}]}]
    responses = [{"question_key": "c1", "choice": c, "screen": 0} for c in (0, 1, 2, 2, 1)] + \
                [{"question_key": "q1", "choice": c, "screen": 1} for c in (0, 1, 1, 1)]
    sig = pc.lesson_signals(screens, responses)
    assert sig == {"not_sure_pct": 60, "quiz_pct": 25, "answers": 9}
    hint = pc.suggestion(14, 22, sig)
    assert hint["carry_on"] and hint["reasons"] == ["you reached slide 14 of 22", "60% weren't sure yet", "only 25% of quiz answers were right"]
    assert pc.suggestion(22, 22, {"not_sure_pct": 10, "quiz_pct": 90}) == {"carry_on": False, "reasons": []}


@test
def run_live_and_present_pick_up_on_the_right_slide():
    screens = [{"slide": 0}, {"slide": 1}, {"slide": 1}, {"slide": 2}, {"slide": 3}]   # slide 2 has two screens
    assert start_screen_for_slide(screens, 3) == 3 and start_screen_for_slide(screens, 2) == 1
    assert start_screen_for_slide([{}, {}, {}], 2) == 1 and start_screen_for_slide([{}], 9) == 0


if __name__ == "__main__":
    failed = 0
    for fn in TESTS:
        try:
            fn()
            print(f"PASS  {fn.__name__}")
        except Exception as e:
            failed += 1
            print(f"FAIL  {fn.__name__}: {e!r}")
    if failed:
        print(f"\n{failed} OF {len(TESTS)} TESTS FAILED")
        sys.exit(1)
    print(f"\nALL {len(TESTS)} TESTS PASSED")
