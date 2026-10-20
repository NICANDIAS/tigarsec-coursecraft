"""
Tests that "this week" and the planning page follow the college calendar in a scheme-of-work
template - induction weeks and half term included - and agree with the printed scheme of work.
Builds a small college-style template (a Word file) in /tmp, so no real files are touched.

Run (from ~/projects/coursecraft):
  docker compose run --rm --no-deps -v "$PWD":/app --entrypoint python app tools/test_scheme_calendar.py
Expected last line: "ALL 3 TESTS PASSED"
"""
import datetime
import os
import sys
import tempfile
from pathlib import Path
from types import SimpleNamespace

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
from docx import Document  # noqa: E402
from app.sow_docx_fill import template_calendar, fill_sow_template_docx  # noqa: E402
from app.sow_calendar import scheme_calendar, scheme_now  # noqa: E402
from app.term_plan import plan_rows, quick_pick  # noqa: E402

TESTS = []
D = lambda m, d: datetime.date(2026, m, d)  # noqa: E731
TMP = tempfile.mkdtemp(prefix="cc_calendar_")

# A college calendar like a real one: two induction/enrichment weeks, then teaching, half term in October.
CALENDAR = [("1-9-2026", "Induction Week"), ("8-9-2026", "Enrichment - Discover"), ("15-9-2026", ""), ("22-9-2026", ""),
            ("29-9-2026", ""), ("6-10-2026", ""), ("13-10-2026", ""), ("20-10-2026", ""), ("27-10-2026", "HALF TERM"),
            ("3-11-2026", ""), ("10-11-2026", "")]


def test(fn):
    TESTS.append(fn)
    return fn


def make_template():
    doc = Document()
    table = doc.add_table(rows=1, cols=4)
    for cell, text in zip(table.rows[0].cells, ["Week commencing", "Topic", "Content", "Assessment"]):
        cell.text = text
    for date, label in CALENDAR:
        row = table.add_row().cells
        row[0].text = date
        row[1].text = label
    path = os.path.join(TMP, "college_template.docx")
    doc.save(path)
    return path


PATH = make_template()
WEEKS = [{"week": n, "title": f"Topic {n}", "learning_outcomes": []} for n in range(1, 9)]
SOW = SimpleNamespace(format_sample_file_path=PATH, append_eligible_labels=None, weeks_json=WEEKS,
                      start_date=D(9, 1), week_anchor_json=None)


@test
def the_template_is_read_row_by_row_with_half_term_left_alone():
    rows = template_calendar(PATH)
    assert [r["available"] for r in rows] == [False, False, True, True, True, True, True, True, False, True, True]
    assert rows[8] == {"date": D(10, 27), "available": False, "label": "HALF TERM"}


@test
def the_planning_page_dates_match_the_printed_scheme():
    cal = scheme_calendar(SOW)
    assert cal["week_dates"][1] == D(9, 15) and cal["week_dates"][6] == D(10, 20) and cal["week_dates"][7] == D(11, 3)
    assert cal["breaks"] == [{"date": D(10, 27), "label": "HALF TERM", "after_week": 6}]
    out = os.path.join(TMP, "printed.docx")
    fill_sow_template_docx(PATH, WEEKS, {}, out)
    printed = {}
    for row in Document(out).tables[0].rows[1:]:
        if row.cells[1].text.startswith("Topic"):
            printed[int(row.cells[1].text.split()[1].split("\n")[0])] = row.cells[0].text
    assert printed[6] == "20-10-2026" and printed[7] == "3-11-2026", printed  # week 7 is after half term in print too


@test
def this_week_and_to_half_term_follow_the_college_calendar():
    now, cal = scheme_now(SOW, None, D(10, 21))         # counting from 1 Sep would say week 8
    assert now["state"] == "teaching" and now["week"] == 6 and now["source"] == "calendar"
    rows = plan_rows(WEEKS, now["week"], {1, 2, 3, 4, 5}, cal["week_dates"], cal["breaks"], D(10, 21))
    assert [r["week"] for r in rows] == [1, 2, 3, 4, 5, 6, None, 7, 8] and rows[6]["label"] == "HALF TERM"
    assert quick_pick(rows, 6, "half") == [6] and quick_pick(rows, 6, "rest") == [6, 7, 8]
    brk, _ = scheme_now(SOW, None, D(10, 28))
    assert brk["state"] == "break" and brk["week"] == 7


if __name__ == "__main__":
    failed = 0
    for fn in TESTS:
        try:
            fn()
            print(f"PASS  {fn.__name__}")
        except Exception as e:
            failed += 1
            print(f"FAIL  {fn.__name__}: {e!r}")
    if failed:
        print(f"\n{failed} OF {len(TESTS)} TESTS FAILED")
        sys.exit(1)
    print(f"\nALL {len(TESTS)} TESTS PASSED")
