"""
Tests for Stage 12 - the tracking sheet, personal revision after a quiz. No AI, no database.
(The reading options on students' phones are checked in a browser - see the patch notes.)

Run (from ~/projects/coursecraft):
  docker compose run --rm --no-deps -v "$PWD":/app --entrypoint python app tools/test_stage12.py
Expected last line: "ALL 6 TESTS PASSED"
"""
import io
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from app import tracker as tr  # noqa: E402
from app import quiz as qz  # noqa: E402
from app.live_activities import _number  # noqa: E402

TESTS = []


def test(fn):
    TESTS.append(fn)
    return fn


UNITS = [
    {"name": "Unit 2 Networking", "learning_outcomes": [{"code": "A", "criteria": [
        {"code": "P1", "description": "Explain topologies."}, {"code": "P10", "description": "Describe a protocol."},
        {"code": "M1", "description": "Compare."}, {"code": "D1", "description": "Evaluate."}]}]},
    {"name": "Unit 5 Security", "learning_outcomes": [{"code": "B", "criteria": [{"code": "P2", "description": "x"}]}]},
    {"name": "Unit 9 Not started", "learning_outcomes": [{"code": "C", "criteria": [{"code": "P3", "description": "y"}]}]},
]
STUDENTS = [{"id": 1, "name": "Amira"}, {"id": 2, "name": "Ben"}, {"id": 3, "name": "Chloe"}]


def work(sid, grade, marks):
    return {"submission_id": 100 + sid, "grade": grade, "marks": [{"criterion_ref": r, "final_mark": m} for r, m in marks]}


MARKS = {
    "Unit 2 Networking": {1: work(1, "Merit", [("P1", "Met"), ("P10", "Met"), ("M1", "Met"), ("D1", "Not met")]),
                          2: work(2, "Not yet achieved", [("P1", "Partly met"), ("P10 Describe", "Met")])},
    "Unit 5 Security": {1: work(1, "Pass", [("P2", "Met")])},
}


@test
def only_units_with_marked_work_are_shown_with_their_criteria():
    t = tr.build_tracker(UNITS, STUDENTS, MARKS)
    assert [u["name"] for u in t["units"]] == ["Unit 2 Networking", "Unit 5 Security"]
    assert t["units"][0]["criteria"] == ["P1", "P10", "M1", "D1"] and t["units"][0]["index"] == 0


@test
def each_cell_is_the_students_result_and_p1_never_mixes_with_p10():
    t = tr.build_tracker(UNITS, STUDENTS, MARKS)
    amira, ben, chloe = t["rows"]
    assert amira["cells"]["Unit 2 Networking"]["crit"] == {"P1": "met", "P10": "met", "M1": "met", "D1": "not"}
    assert ben["cells"]["Unit 2 Networking"]["crit"] == {"P1": "partly", "P10": "met", "M1": "", "D1": ""}
    assert amira["cells"]["Unit 2 Networking"]["grade"] == "Merit" and amira["cells"]["Unit 2 Networking"]["submission_id"] == 101
    assert chloe["cells"]["Unit 2 Networking"]["has_work"] is False


@test
def students_at_risk_are_flagged_with_the_reason():
    t = tr.build_tracker(UNITS, STUDENTS, MARKS)
    amira, ben, chloe = t["rows"]
    assert amira["risk"] == []
    assert ben["risk"] == ["not yet passing Unit 2 Networking"]
    assert chloe["risk"] == ["no marked work for Unit 2 Networking"]  # 2 of 3 handed it in; Unit 5 only 1 of 3 - not flagged
    assert t["at_risk"] == 2
    assert t["unit_summary"]["Unit 2 Networking"]["criteria_met_pct"]["P1"] == 50
    assert t["unit_summary"]["Unit 2 Networking"]["grades"] == {"Not yet achieved": 1, "Merit": 1}


@test
def the_excel_download_matches_the_page():
    from openpyxl import load_workbook
    book = load_workbook(io.BytesIO(tr.tracker_xlsx("T Level Y2", tr.build_tracker(UNITS, STUDENTS, MARKS))))
    ws = book.active
    assert ws["A1"].value == "T Level Y2 - tracking sheet" and ws["B3"].value == "Unit 2 Networking"
    assert [ws.cell(row=4, column=c).value for c in range(2, 7)] == ["P1", "P10", "M1", "D1", "Grade"]
    assert [ws.cell(row=5, column=c).value for c in range(1, 7)] == ["Amira", "✓", "✓", "✓", "✗", "Merit"]
    assert ws.cell(row=5, column=2).fill.fgColor.rgb.endswith("C6EFCE")
    assert ws.cell(row=7, column=ws.max_column).value == "no marked work for Unit 2 Networking"


QUESTIONS = [{"question": "2+2?", "choices": ["3", "4"], "answer": 1, "explanation": "Add them."},
             {"question": "Capital of France?", "choices": ["Paris", "Rome"], "answer": 0, "explanation": ""},
             {"question": "HTTPS port?", "choices": ["80", "443"], "answer": 1, "explanation": "443 is HTTPS."}]


@test
def revision_is_the_questions_a_student_got_wrong_or_missed():
    screens = _number(qz.screens_from_quiz(QUESTIONS, 20, "Mixed"))
    keys = [q["key"] for s in screens for q in s["questions"]]
    rev = qz.revision_for(screens, {keys[0]: 1, keys[1]: 1})  # right, wrong, not answered
    assert [r["question"] for r in rev] == ["Capital of France?", "HTTPS port?"]
    assert rev[0]["picked"] == 1 and rev[0]["answer"] == 0 and rev[1]["picked"] is None
    assert rev[1]["explanation"] == "443 is HTTPS."


@test
def a_perfect_quiz_means_nothing_to_revise():
    screens = _number(qz.screens_from_quiz(QUESTIONS, 20, "Mixed"))
    keys = [q["key"] for s in screens for q in s["questions"]]
    assert qz.revision_for(screens, dict(zip(keys, [1, 0, 1]))) == []


if __name__ == "__main__":
    failed = 0
    for fn in TESTS:
        try:
            fn()
            print(f"PASS  {fn.__name__}")
        except AssertionError as e:
            failed += 1
            print(f"FAIL  {fn.__name__}: {e}")
    if failed:
        print(f"\n{failed} OF {len(TESTS)} TESTS FAILED")
        sys.exit(1)
    print(f"\nALL {len(TESTS)} TESTS PASSED")
