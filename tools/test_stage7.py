"""
Tests for Stage 7 - presenting in the browser exactly as the downloaded PowerPoint. No AI,
no database. Builds real .pptx files with CourseCraft's own generator, then checks that the
browser view is read from THAT file (positions, colours, fonts, text, pictures) and that live
lessons line up with its slides.

Run (from ~/projects/coursecraft):
  docker compose run --rm --no-deps -v "$PWD":/app --entrypoint python app tools/test_stage7.py
Expected last line: "ALL 8 TESTS PASSED"
"""
import os
import random
import sys
import tempfile
import types
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
_stub = types.ModuleType("app.llm_gateway")
_stub.complete_json = lambda **k: (_ for _ in ()).throw(RuntimeError("AI disabled in tests"))
_stub.get_effective_provider = lambda *a, **k: "none"
sys.modules["app.llm_gateway"] = _stub

from pptx import Presentation  # noqa: E402
from app import pptx_generator as pg  # noqa: E402
from app import pptx_html as ph  # noqa: E402
from app import live_activities as la  # noqa: E402

TESTS = []
TMP = tempfile.mkdtemp(prefix="cc_stage7_")


def test(fn):
    TESTS.append(fn)
    return fn


def photo(name, size):
    from PIL import Image
    path = os.path.join(TMP, name)
    Image.new("RGB", size, (40, 90, 140)).save(path)
    return path


OUTLINE = {"lesson_title": "Cash flow forecasting", "slides": [
    {"type": "title", "title": "Cash flow forecasting", "subtitle": "Component 1"},
    {"type": "retrieval_starter", "title": "Do Now", "questions": ["What is revenue?", "Name a fixed cost."], "time_minutes": 5},
    {"type": "objectives", "objectives": ["Explain cash flow", "Construct a forecast"]},
    {"type": "content", "title": "What is cash flow?", "guiding_question": "Can a profitable firm run out of cash?",
     "key_points": ["Money in and out over time", "Timing matters"], "check_question": "Why does timing matter?", "image_query": "till",
     "resolved_image": {"type": "photo", "path": "", "attribution": "Photo by Sam on Unsplash"}},
    {"type": "content", "title": "Balances", "key_points": ["Closing = opening + net"], "check_question": "What does negative mean?",
     "image_query": "bank", "resolved_image": {"type": "generated"}},
    {"type": "real_life_example", "title": "A cafe", "scenario": "Busy summers.", "discussion_question": "Winter?", "image_query": "cafe",
     "resolved_image": {"type": "generated"}},
    {"type": "activity", "title": "Build a forecast", "instructions": ["Use the sheet", "Find net cash flow"], "grouping": "pairs", "time_minutes": 15},
    {"type": "practical", "title": "Spreadsheet", "instructions": ["Open the template"], "materials_needed": "Laptops", "time_minutes": 20},
    {"type": "retrieval_plenary", "title": "Check", "questions": ["Define net cash flow"]},
    {"type": "objectives_review", "objectives_check": [{"objective": "Explain cash flow", "self_assess_prompt": "Could I explain it?"}]},
    {"type": "something_new", "title": "Not a slide type the generator knows"},
]}


def build(outline=OUTLINE, seed=1):
    random.seed(seed)
    outline = {**outline, "slides": [dict(s) for s in outline["slides"]]}
    for s in outline["slides"]:
        if s.get("resolved_image", {}).get("type") == "photo":
            s["resolved_image"] = {**s["resolved_image"], "path": photo("p.jpg", (1200, 800))}
    path = os.path.join(TMP, f"deck_{seed}.pptx")
    pg.render_lesson_pptx(outline, None, TMP, path)
    return outline, path


@test
def every_slide_of_the_file_is_read_with_its_background():
    outline, path = build()
    layout = ph.read_slides(path)
    prs = Presentation(path)
    assert round(layout["width"]) == 1280 and round(layout["height"]) == 720  # 13.333" x 7.5" at 96 px per inch
    assert len(layout["slides"]) == len(prs.slides) == 11  # the content slide with a guiding question is TWO slides
    palette = pg.pick_palette("Cash flow forecasting")
    assert layout["slides"][0]["background"] == palette["accent"]           # title slide: accent colour
    assert layout["slides"][7]["background"] == palette["tint"]             # activity (8th slide): tinted
    assert layout["slides"][1]["background"] == "FFFFFF"


@test
def text_keeps_its_exact_position_size_colour_and_font():
    outline, path = build()
    title = ph.read_slides(path)["slides"][1]["elements"][0]
    assert (title["x"], title["y"], title["w"], title["h"]) == (57.6, 48.0, 1164.77, 86.4)  # 0.6", 0.5", 12.13" x 0.9"
    run = title["text"]["paragraphs"][0]["runs"][0]
    assert run == {"text": "Do Now", "size": 32.0, "bold": True, "italic": False,
                   "color": pg.pick_palette("Cash flow forecasting")["accent"], "font": "Calibri"}
    assert title["text"]["insets"] == [9.6, 4.8, 9.6, 4.8] and title["text"]["anchor"] == "top"
    bullets = ph.read_slides(path)["slides"][1]["elements"][2]["text"]["paragraphs"]
    assert [p["runs"][0]["text"] for p in bullets] == ["•  What is revenue?", "•  Name a fixed cost."] and bullets[0]["space_after"] == 16.0


@test
def badges_and_placeholder_graphics_are_shapes_with_their_colours():
    outline, path = build()
    badge = ph.read_slides(path)["slides"][1]["elements"][1]
    assert badge["kind"] == "shape" and badge["geometry"] == "roundRect" and badge["shadow"] is True
    assert badge["text"]["anchor"] == "middle" and badge["text"]["paragraphs"][0]["align"] == "center"
    assert badge["text"]["paragraphs"][0]["runs"][0]["font"] == "Calibri"  # no font set -> the theme font, as PowerPoint shows
    shapes = [e for e in ph.read_slides(path)["slides"][5]["elements"] if e["kind"] == "shape"]  # the generated image card
    assert {e["geometry"] for e in shapes} == {"roundRect", "ellipse"}


@test
def pictures_come_out_of_the_file_itself():
    outline, path = build()
    layout = ph.read_slides(path)
    pictures = [(i, e) for i, s in enumerate(layout["slides"]) for e in s["elements"] if e["kind"] == "picture"]
    assert len(pictures) == 2  # question + answer slide of the first content entry
    s_index, shape_index = map(int, pictures[0][1]["media"].split("/"))
    blob, content_type = ph.picture_bytes(path, s_index, shape_index)
    assert content_type == "image/jpeg" and blob[:2] == b"\xff\xd8"
    assert ph.picture_bytes(path, 1, 0) is None and ph.picture_bytes(path, 99, 0) is None  # not a picture / no such slide


@test
def the_image_side_matches_the_file_even_though_it_is_random():
    sides = set()
    for seed in range(6):
        outline, path = build(seed=seed)
        pic = next(e for e in ph.read_slides(path)["slides"][3]["elements"] if e["kind"] == "picture")
        sides.add("left" if pic["x"] < 300 else "right")
    assert sides == {"left", "right"}  # the generator picks a side at random - the browser shows whichever the FILE has


@test
def the_slide_map_matches_what_the_generator_really_produces():
    for seed in range(5):
        rng = random.Random(seed)
        slides = [rng.choice(OUTLINE["slides"][:10]) for _ in range(rng.randint(3, 14))]
        outline, path = build({"lesson_title": "x", "slides": slides}, seed=100 + seed)
        assert sum(len(m) for m in ph.outline_slide_map(outline)) == len(Presentation(path).slides)


@test
def live_screens_line_up_with_the_powerpoint_slides():
    outline, path = build()
    screens = la.screens_from_deck(outline)
    count = len(ph.read_slides(path)["slides"])
    known = [s for s in screens]
    screens = la.screens_from_deck({**outline, "slides": outline["slides"][:10]})
    assert la.attach_pptx_slides(screens, {**outline, "slides": outline["slides"][:10]}, count) is True
    assert [s["slide"] for s in screens] == list(range(11))
    q, a = screens[3], screens[4]
    assert q["kind"] == "content_question" and q["questions"][0]["prompt"] == "Can a profitable firm run out of cash?" and not q["body"]
    assert a["kind"] == "content" and a["body"] == ["Money in and out over time", "Timing matters"]
    assert known[-1]["title"].startswith("Not a slide type")  # an unknown slide type still shows as text...
    assert la.attach_pptx_slides(known, outline, count) is False  # ...so it can't be paired with the file: text view is kept


@test
def a_shape_added_by_hand_in_powerpoint_is_still_drawn():
    from pptx.util import Inches
    from pptx.enum.shapes import MSO_SHAPE
    outline, path = build()
    prs = Presentation(path)
    star = prs.slides[1].shapes.add_shape(MSO_SHAPE.STAR_5_POINT, Inches(1), Inches(5), Inches(1), Inches(1))
    star.text = "Well done"
    edited = os.path.join(TMP, "edited.pptx")
    prs.save(edited)
    added = ph.read_slides(edited)["slides"][1]["elements"][-1]
    assert added["kind"] == "shape" and added["geometry"] == "rect" and added["text"]["paragraphs"][0]["runs"][0]["text"] == "Well done"


if __name__ == "__main__":
    failed = 0
    for fn in TESTS:
        try:
            fn()
            print(f"PASS  {fn.__name__}")
        except AssertionError as e:
            failed += 1
            print(f"FAIL  {fn.__name__}: {e}")
    if failed:
        print(f"\n{failed} OF {len(TESTS)} TESTS FAILED")
        sys.exit(1)
    print(f"\nALL {len(TESTS)} TESTS PASSED")
